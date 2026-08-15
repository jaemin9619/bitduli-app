from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = r"C:\Users\명재민\Documents\26빅데이터\삐뚤_제안서_현재구조_재작성본.docx"
FONT = "Malgun Gothic"


def set_run(run, size=11, bold=None, color=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for key in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rfonts.set(qn(key), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(p, size=11, color=None):
    for run in p.runs:
        set_run(run, size=size, color=color)


def set_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, val in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[idx]))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_margins(cell)


def setup_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    return doc


def para(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    style_paragraph(p)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.194)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)
        style_paragraph(p)


def numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.45)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)
        style_paragraph(p)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        set_shading(c, "F4F6F9")
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_paragraph(p, size=10, color="0B2545")
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                style_paragraph(p, size=10)
    set_table_geometry(t, widths)
    doc.add_paragraph()
    return t


doc = setup_doc()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("삐뚤: AI 그림일기 서비스 제안서")
set_run(r, size=24, bold=True, color="0B2545")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("현재 26빅데이터 프로젝트 내부 구조 반영 재작성본")
set_run(r, size=12, color="555555")

para(
    doc,
    "본 문서는 현재 26빅데이터 폴더의 `bitduli-app` 구조를 기준으로 제안서를 다시 작성한 것이다. "
    "Claude와 함께 수정된 뒤 앱의 내부 구조는 ‘React 모바일 웹앱 + Firebase Auth/Firestore + Firebase Functions AI 처리 + PWA 배포’로 정리되어 있다. "
    "초기 제안서의 Google Calendar 연동, 데스크톱 중심 화면, 복잡한 캐릭터 외형 설정은 현재 구현 범위와 맞지 않으므로 제외하거나 향후 확장 과제로 이동했다.",
)

doc.add_heading("1. 프로젝트 개요", level=1)
doc.add_heading("1.1 서비스명과 한 줄 정의", level=2)
para(doc, "서비스명: 삐뚤")
para(doc, "정의: 사용자가 짧은 하루 일기를 쓰면 AI가 아이 말투의 요약문과 삐뚤빼뚤한 크레용 그림일기로 바꿔 주는 모바일 우선 웹앱이다.")

doc.add_heading("1.2 현재 구현 기준의 문제 인식", level=2)
para(
    doc,
    "기존 일기 서비스는 긴 글을 직접 쓰고 사진을 첨부하는 방식에 머무는 경우가 많다. 삐뚤은 일기 작성의 부담을 낮추기 위해 사용자가 사건과 감정을 짧게 입력하면, "
    "AI가 이를 그림일기 형태로 재구성한다. 현재 앱은 완벽한 이미지보다 아이가 그린 듯한 서툰 선, 감정 색상, 날씨·날짜 맥락을 결합해 ‘다시 보고 싶은 기록’을 만드는 데 초점을 둔다.",
)

doc.add_heading("1.3 현재 구조에서 확정된 범위", level=2)
bullets(
    doc,
    [
        "서비스 실행 단위는 `bitduli-app`이며, 이전 `lineme-diary`와 `_bbiddul_zip_restore`는 참고·복원본 성격이다.",
        "앱 진입 흐름은 Splash, Google Login, Customization, CalendarDashboard 네 단계로 구성된다.",
        "사용자 데이터는 `users/{email}` 문서와 하위 `diaries`, `friends` 서브컬렉션에 저장된다.",
        "AI 요약과 이미지 생성은 Firebase Functions callable 함수가 담당한다.",
        "친구 요청·승인·삭제와 일기 저장·공개 토글은 현재 대시보드에서 Firestore를 직접 갱신하는 구조다.",
        "생성 이미지는 현재 클라이언트에서 압축한 data URL 형태로 일기 문서의 `imageUrl`에 저장된다.",
    ],
)

doc.add_heading("2. 현재 앱 내부 구조", level=1)
doc.add_heading("2.1 폴더와 역할", level=2)
table(
    doc,
    ["경로", "역할", "제안서 반영 방식"],
    [
        ["bitduli-app/src", "React 화면과 상태 관리", "사용자 플로우와 UI/UX 구현 범위의 기준"],
        ["bitduli-app/src/components", "Splash, Login, Customization, CalendarDashboard, Logo", "주요 화면 설명의 기준"],
        ["bitduli-app/src/lib", "Firebase 초기화와 Functions callable 래퍼", "클라이언트-백엔드 연결 구조 설명"],
        ["bitduli-app/functions", "Gemini/Imagen 호출, 친구 callable 함수 정의", "AI 처리와 서버리스 구성 설명"],
        ["bitduli-app/firestore.rules", "사용자·일기·친구 접근 권한", "공개 일기와 친구 권한 모델 설명"],
        ["bitduli-app/public", "PWA manifest, service worker, 아이콘", "설치형 웹앱과 캐싱 설명"],
    ],
    [2450, 3300, 3610],
)

doc.add_heading("2.2 화면 흐름", level=2)
numbered(
    doc,
    [
        "SplashView: 노란 배경과 삐뚤 로고로 시작 화면을 제공한다.",
        "LoginView: Firebase Authentication의 Google Provider로 로그인한다.",
        "CustomizationView: 최초 사용자의 닉네임과 대표 색상을 받아 Firestore 사용자 프로필을 만든다.",
        "CalendarDashboard: 캘린더, 그림책, 쓰기, 친구, 설정을 하단 탭으로 제공하는 실제 메인 앱이다.",
    ],
)

doc.add_heading("2.3 대시보드 내부 상태", level=2)
para(
    doc,
    "CalendarDashboard는 현재 앱에서 가장 큰 상태 관리 단위다. 월간 날짜, 선택 날짜, 작성 중인 일기, 생성 이미지, 생성 요약, 공개 여부, 친구 목록, 친구 요청 상태, 설정 하위 화면을 한 컴포넌트 안에서 관리한다. "
    "따라서 제안서의 기능 구조도 단순 ‘AI 그림 생성 앱’이 아니라 ‘일기 기록, 회상, 친구 공개, 계정 관리가 결합된 모바일 일기장’으로 표현하는 것이 맞다.",
)
table(
    doc,
    ["상태/데이터", "현재 의미"],
    [
        ["diaryEntries", "로그인 사용자의 날짜별 일기 목록. Firestore와 localStorage 캐시를 함께 사용"],
        ["activeTab", "calendar, book, write, view, friends, settings 화면 전환"],
        ["diaryText, diaryTitle, diaryWeather, selectedFeeling", "일기 작성 입력값"],
        ["generatedImageUrl, generatedSummary", "AI 생성 결과와 요약문"],
        ["isWritePublic, viewingEntry.isPublic", "친구에게 공개할지 결정하는 공개 상태"],
        ["friendsList, pendingSentRequests, pendingReceivedRequests", "친구 관계와 요청 상태"],
        ["settingsTab", "프로필 수정, 친구 초대·승인·삭제, 계정 삭제 화면"],
    ],
    [3150, 6210],
)

doc.add_heading("3. 핵심 기능 재정의", level=1)
doc.add_heading("3.1 일기 작성과 AI 생성", level=2)
para(
    doc,
    "사용자는 제목, 본문, 날씨, 감정을 입력한 뒤 AI 그림 생성을 실행한다. 클라이언트는 `generateDiaryDrawing` callable function을 호출하고, 반환된 이미지가 PNG base64이면 브라우저 canvas에서 축소·JPEG 압축한 뒤 data URL로 저장한다. "
    "동시에 `summarizeDiary` callable function을 호출해 아이 말투의 짧은 요약문을 만든다. API 실패나 키 미설정 상황에서는 SVG 기반 fallback이 동작하도록 설계되어 있다.",
)

doc.add_heading("3.2 캘린더와 그림책", level=2)
para(
    doc,
    "홈 탭은 월간 캘린더를 중심으로 한다. 작성된 날짜는 감정 색상과 미니 얼굴 로고로 표시되어, 사용자가 한 달의 감정 흐름을 한눈에 볼 수 있다. "
    "그림책 탭은 같은 일기 데이터를 월간 책 형태로 다시 보여 주며, AI 이미지와 요약문 중심의 회상 경험을 제공한다.",
)

doc.add_heading("3.3 친구 공개 구조", level=2)
para(
    doc,
    "친구 기능은 이메일 기반 요청, 수락, 거절, 삭제로 구성된다. 현재 UI는 Firestore의 양쪽 사용자 `friends` 문서를 직접 생성·수정·삭제한다. "
    "친구로 수락된 사용자만 상대의 공개 일기를 읽을 수 있고, 이 조건은 `firestore.rules`의 `isAcceptedFriend` 함수와 `resource.data.isPublic == true` 조건으로 보호된다.",
)

doc.add_heading("3.4 설정과 계정 관리", level=2)
para(
    doc,
    "설정 탭에는 프로필 수정, 친구 설정, 로그아웃, 계정 삭제 기능이 포함된다. 계정 삭제는 사용자의 일기 문서, 친구 관계 문서, 사용자 프로필 문서를 삭제하고 localStorage 캐시를 정리하는 흐름으로 구현되어 있다. "
    "이 기능은 초기 제안서에는 약했지만 현재 앱에서는 실제 서비스 운영 관점에서 중요한 내부 구조가 되었다.",
)

doc.add_heading("4. 기술 구성", level=1)
table(
    doc,
    ["영역", "현재 기술/파일", "역할"],
    [
        ["Frontend", "React 19, TypeScript, Vite, Tailwind CSS, motion/react, lucide-react", "모바일 우선 UI와 화면 전환"],
        ["App shell", "App.tsx, SplashView, LoginView, CustomizationView, CalendarDashboard", "인증 상태에 따른 진입 흐름과 메인 기능 제공"],
        ["Auth", "Firebase Authentication, GoogleAuthProvider", "Google 계정 로그인과 이메일 기반 사용자 식별"],
        ["Database", "Cloud Firestore", "프로필, 일기, 친구 상태, AI 사용 로그 저장"],
        ["Functions", "Firebase Functions v2, asia-northeast3, Node 22", "AI 요약·이미지 생성, callable 백엔드 함수"],
        ["AI", "Gemini 2.5 Flash, Imagen 4.0 Generate", "요약문 생성, 장면 설명 추출, 그림일기 이미지 생성"],
        ["PWA/Hosting", "Firebase Hosting, manifest.webmanifest, sw.js", "배포, 앱 설치형 실행, 앱 셸 캐싱"],
        ["Security", "firestore.rules, storage.rules", "본인 데이터 보호, 친구 공개 읽기 제한, Storage 쓰기 차단"],
    ],
    [1800, 3400, 4160],
)

doc.add_heading("4.1 데이터 모델", level=2)
table(
    doc,
    ["문서 경로", "주요 필드", "현재 의미"],
    [
        ["users/{email}", "nickname, email, profileColorName, profileColorHex, createdAt, updatedAt", "사용자 프로필과 대표 색상"],
        ["users/{email}/diaries/{date}", "date, title, text, imageUrl, feeling, weather, isPublic, createdAt, summary", "날짜별 그림일기 원문·그림·요약·공개 상태"],
        ["users/{email}/friends/{friendEmail}", "friendEmail, friendNickname, friendColorHex, status, updatedAt", "sent, received, accepted 기반 친구 관계"],
        ["aiUsage/{docId}", "ownerEmail, type, createdAt", "AI 요약·이미지 생성 로그. 클라이언트 직접 접근 차단"],
    ],
    [2450, 3900, 3010],
)

doc.add_heading("4.2 보안 규칙 기준", level=2)
bullets(
    doc,
    [
        "사용자 문서는 로그인 이메일과 문서 ID가 일치하는 사용자만 생성·수정·삭제할 수 있다.",
        "일기 문서는 본인만 쓸 수 있고, 읽기는 본인이거나 ‘수락된 친구 + 공개 일기’ 조건일 때만 허용된다.",
        "친구 문서는 본인이 관리하되, 상대 사용자 문서에 자기 자신을 나타내는 친구 문서를 생성·수정·삭제할 수 있도록 허용한다.",
        "Storage는 `diary-images/{userId}/{fileName}` 읽기만 인증 사용자에게 열려 있고 쓰기는 차단되어 있다. 현재 이미지 저장은 Storage가 아니라 Firestore data URL 중심이다.",
        "AI 사용 로그 `aiUsage`는 클라이언트 read/write를 모두 차단해 백엔드 기록용으로만 둔다.",
    ],
)

doc.add_heading("5. 초기 제안서 대비 수정된 방향", level=1)
table(
    doc,
    ["항목", "초기 제안서", "현재 구조 기준 제안서"],
    [
        ["서비스명", "LineMe Diary", "삐뚤"],
        ["핵심 이미지 스타일", "라인 아트 중심", "아이의 크레용 질감, 삐뚤한 선, 색 번짐"],
        ["메인 플랫폼", "데스크톱 또는 범용 웹", "세로형 모바일 우선 PWA"],
        ["일정 연동", "Google Calendar API", "앱 내부 월간 캘린더와 일기 날짜 관리"],
        ["캐릭터 설정", "머리·의상·액세서리", "닉네임과 대표 색상 중심"],
        ["AI 백엔드", "Vertex AI 중심 설명", "Firebase Functions에서 Gemini/Imagen 호출"],
        ["친구 기능", "소셜 피드, 좋아요, 댓글 중심", "친구 요청·수락과 공개 일기 조회 중심"],
        ["이미지 저장", "Cloud Storage URL 전제", "현재는 압축 data URL을 Firestore에 저장, Storage 전환은 개선 과제"],
        ["운영 기능", "상대적으로 약함", "프로필 수정, 로그아웃, 계정 삭제, 보안 규칙 포함"],
    ],
    [1700, 3200, 4460],
)

doc.add_heading("6. 현재 구조에서 남은 개선 과제", level=1)
bullets(
    doc,
    [
        "일부 UI 라벨과 샘플 데이터에 남아 있는 인코딩 깨짐을 정리해야 한다.",
        "CalendarDashboard가 너무 많은 상태와 화면을 한 컴포넌트에서 담당하므로, 화면 단위와 데이터 훅 단위로 분리하면 유지보수가 쉬워진다.",
        "이미지를 Firestore 문서에 data URL로 넣는 방식은 빠른 구현에는 좋지만 문서 크기 제한 리스크가 있으므로, Cloud Storage 저장 후 URL만 보관하는 구조가 적합하다.",
        "src/lib/backend.ts에 친구 callable 래퍼가 있지만 현재 대시보드는 Firestore 직접 쓰기를 사용한다. 한쪽 방식으로 정리하면 구조가 더 명확해진다.",
        "README가 AI Studio 기본 문서 상태이므로 실제 삐뚤 실행·배포·환경변수 안내로 교체해야 한다.",
        "아동 대상 서비스 성격을 고려해 신고, 차단, 보호자 확인, 공개 범위 세분화 같은 안전 기능을 추가할 수 있다.",
        "공유 시트는 현재 시뮬레이션 성격이 강하므로 Web Share API, 이미지 다운로드, 월간 그림책 내보내기로 확장할 수 있다.",
    ],
)

doc.add_heading("7. 기대 효과", level=1)
doc.add_heading("7.1 기록 부담 완화", level=2)
para(doc, "짧은 텍스트, 감정, 날씨만 입력해도 AI가 그림과 요약문을 만들어 주기 때문에 사용자는 완성도 부담 없이 하루를 기록할 수 있다.")
doc.add_heading("7.2 회상 경험 강화", level=2)
para(doc, "캘린더와 그림책 화면은 날짜별 기록을 단순 목록이 아니라 감정 색상과 그림 중심의 회상 경험으로 바꾼다.")
doc.add_heading("7.3 제한된 관계 안의 공유", level=2)
para(doc, "친구로 수락된 사용자에게만 공개 일기를 보여 주므로, 일반 SNS보다 작은 관계 안에서 부담 없이 일기를 공유할 수 있다.")
doc.add_heading("7.4 실제 서비스 구조 학습", level=2)
para(doc, "삐뚤은 프론트엔드 화면, 인증, Firestore 데이터 모델, 보안 규칙, 서버리스 AI 처리, PWA 배포까지 포함하므로 단순 디자인 시안보다 실제 배포 가능한 AI 앱 구조를 보여 준다.")

doc.add_heading("8. 결론", level=1)
para(
    doc,
    "현재 26빅데이터 프로젝트의 삐뚤은 초기 제안서보다 내부 구조가 구체화되었다. 서비스의 중심은 더 이상 Google Calendar나 데스크톱형 생산성 도구가 아니라, "
    "모바일에서 아이답게 하루를 쓰고 AI 그림으로 다시 보는 그림일기 경험이다. 제안서도 이 구조에 맞춰 React/Firebase 기반 앱, Firestore 데이터 모델, Functions AI 파이프라인, "
    "친구 공개 권한, PWA 배포를 핵심으로 설명해야 한다.",
)
para(
    doc,
    "따라서 향후 개발은 기능을 더 크게 늘리기보다, 현재 구현된 구조를 안정화하는 방향이 우선이다. 특히 인코딩 정리, 컴포넌트 분리, 이미지 저장 구조 개선, README와 배포 문서 정비, "
    "아동 서비스 안전 기능 보강을 진행하면 제안서와 실제 앱의 일치도가 높아진다.",
)

footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run("삐뚤 제안서 현재구조 재작성본")
set_run(r, size=9, color="555555")

doc.save(OUT)
print(OUT)
