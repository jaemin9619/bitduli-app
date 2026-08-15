from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = r"C:\Users\명재민\Documents\26빅데이터\삐뚤_제안서_수정본.docx"


def set_run_font(run, name="Malgun Gothic", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_font(paragraph, name="Malgun Gothic", size=11, color=None):
    for run in paragraph.runs:
        set_run_font(run, name=name, size=size, color=color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[idx]))
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Malgun Gothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run("삐뚤: AI 캐릭터와 함께하는 그림일기 서비스")
    set_run_font(run, size=24, bold=True, color="0B2545")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("현재 구현 기능 기준 개정 제안서")
    set_run_font(run, size=12, color="555555")


def add_para(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    set_paragraph_font(p)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.194)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.208
        p.add_run(item)
        set_paragraph_font(p)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = text
        set_cell_shading(hdr[i], "F4F6F9")
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_font(p, size=10, color="0B2545")
            for run in p.runs:
                run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                set_paragraph_font(p, size=10)
    set_table_widths(table, widths)
    doc.add_paragraph()
    return table


doc = Document()
style_document(doc)
add_title(doc)

add_para(
    doc,
    "본 문서는 초기 제안서의 서비스명, 기능 범위, 기술 구성, 개발 로드맵을 현재 구현된 앱 ‘삐뚤’ 기준으로 개정한 것이다. "
    "초기안의 Google Calendar 연동, 데스크톱 중심 Split View, 복잡한 외형 커스터마이징은 현재 앱 범위에서 제외하고, "
    "실제 구현된 모바일 우선 그림일기 작성, AI 요약·이미지 생성, 친구 공개 일기 공유, PWA 배포 흐름을 중심으로 재정리했다.",
)

doc.add_heading("1. 프로젝트 개요", level=1)
doc.add_heading("1.1 프로젝트명", level=2)
add_para(doc, "삐뚤")
add_para(doc, "부제: AI가 도와주는 아이 그림일기 앱")

doc.add_heading("1.2 프로젝트 배경", level=2)
add_para(
    doc,
    "아이들이 하루를 글과 그림으로 기록하는 그림일기는 감정 표현, 회상, 자기 이해를 돕는 활동이다. "
    "하지만 직접 그림을 그리거나 긴 글을 쓰는 과정은 부담이 될 수 있고, 보호자 입장에서도 꾸준히 기록을 유지하기 어렵다. "
    "삐뚤은 사용자가 짧은 일기를 쓰면 AI가 내용을 이해해 아이가 그린 듯한 삐뚤빼뚤한 그림과 짧은 그림일기 문장으로 바꿔 주는 서비스다.",
)
add_para(
    doc,
    "현재 구현된 앱은 ‘완벽하게 정돈된 기록’보다 ‘아이답고 서툰 기록의 매력’을 중심에 둔다. "
    "감정 색상, 날씨, 일기 문장, AI 그림이 한 장의 그림일기 카드로 결합되며, 사용자는 달력과 그림책 화면에서 자신의 기록을 다시 볼 수 있다.",
)

doc.add_heading("1.3 기존 제안 대비 개정 사항", level=2)
add_table(
    doc,
    ["항목", "초기 제안서", "삐뚤 현재 기준"],
    [
        ["서비스명", "LineMe Diary", "삐뚤"],
        ["핵심 콘셉트", "라인 아트 중심 AI 그림일기", "아이의 크레용 질감과 삐뚤한 선을 살린 AI 그림일기"],
        ["주요 화면", "데스크톱 최적화, Navigation Rail, Split View", "모바일 우선 앱 셸, 하단 탭 네비게이션, 캘린더·쓰기·그림책·친구·설정"],
        ["일정 연동", "Google Calendar API 연동", "현재 범위에서는 제외, 앱 내부 캘린더로 일기 기록 관리"],
        ["커스터마이징", "머리·의상·액세서리 등 캐릭터 외형 설정", "닉네임과 대표 색상 중심의 간단한 프로필 설정"],
        ["AI 구성", "Vertex AI 중심 설명", "Firebase Functions에서 Gemini 2.5 Flash와 Imagen 4 호출"],
        ["저장 방식", "Cloud Storage URL 중심", "Firestore 일기 문서에 압축 이미지 데이터와 메타데이터 저장"],
    ],
    [1700, 3300, 4360],
)

doc.add_heading("1.4 프로젝트 목표", level=2)
add_bullets(
    doc,
    [
        "짧은 글 입력만으로 아이답고 귀여운 그림일기 결과물을 생성한다.",
        "날씨와 감정을 함께 기록해 하루의 분위기를 쉽게 회상할 수 있게 한다.",
        "캘린더와 월간 그림책 화면으로 기록을 자연스럽게 다시 보게 한다.",
        "친구 공개 범위를 통해 무분별한 SNS가 아닌 작은 관계 안에서 일기를 공유한다.",
        "Firebase 기반 인증·저장·배포 구조로 실제 사용 가능한 웹앱 형태를 완성한다.",
    ],
)

doc.add_heading("2. 주요 기능", level=1)
doc.add_heading("2.1 Google 로그인과 프로필 설정", level=2)
add_para(
    doc,
    "사용자는 Google 계정으로 로그인하며, 최초 로그인 후 Firestore에 사용자 프로필을 생성한다. "
    "프로필에는 닉네임, 이메일, 대표 색상 정보가 저장된다. 대표 색상은 삐뚤 로고와 일기 화면의 시각적 분위기에 반영되어 사용자별 식별성을 만든다.",
)

doc.add_heading("2.2 일기 작성", level=2)
add_para(
    doc,
    "쓰기 화면에서는 제목, 본문, 날씨, 감정을 입력한다. 감정은 행복, 피곤, 평온, 슬픔, 화남 등으로 제공되며 각 감정은 전용 색상과 표정 아이콘으로 표현된다. "
    "작성자는 일기 공개 여부를 선택할 수 있고, 공개된 일기만 친구에게 표시된다.",
)

doc.add_heading("2.3 AI 요약과 그림 생성", level=2)
add_para(
    doc,
    "Firebase Functions는 사용자의 일기 본문과 감정 정보를 받아 Gemini 2.5 Flash로 아이가 쓴 듯한 짧은 한국어 요약문을 만든다. "
    "이후 Imagen 4에 전달할 장면 설명을 구성해, 실제 일기 내용의 장소·사물·행동·감정을 반영한 정사각형 그림일기 이미지를 생성한다.",
)
add_para(
    doc,
    "이미지 프롬프트는 ‘6~7세 아이가 거친 종이에 크레용으로 그린 그림’이라는 스타일을 기준으로 하며, 색이 선 밖으로 번지고 삐뚤한 형태가 남도록 유도한다. "
    "AI 키가 없거나 이미지 생성에 실패하는 경우에는 SVG 기반 대체 그림을 반환해 기능 흐름이 끊기지 않도록 설계되어 있다.",
)

doc.add_heading("2.4 캘린더와 일기 상세 보기", level=2)
add_para(
    doc,
    "홈 화면은 월간 캘린더를 중심으로 구성된다. 작성된 날짜에는 감정 색상 카드와 미니 표정 로고가 표시되어 어떤 날에 어떤 감정의 일기를 썼는지 한눈에 알 수 있다. "
    "날짜를 선택하면 해당 일기의 그림, 짧은 요약문, 원문, 날씨, 감정, 공개 상태를 확인할 수 있다.",
)

doc.add_heading("2.5 월간 그림책", level=2)
add_para(
    doc,
    "그림책 탭은 한 달 동안 작성한 그림일기를 책처럼 넘겨 보는 화면이다. 사용자는 AI 그림과 한 글자씩 배치된 요약문을 통해 자신의 한 달을 그림책 형태로 회상한다. "
    "초기 제안서의 ‘월간 그림책’ 아이디어는 현재 구현에서 핵심 탐색 화면으로 구체화되었다.",
)

doc.add_heading("2.6 친구와 공개 일기 공유", level=2)
add_para(
    doc,
    "친구 기능은 이메일 기반 요청, 수락, 거절, 삭제 흐름으로 구성된다. 친구 관계가 수락된 사용자만 상대의 공개 일기를 조회할 수 있으며, Firestore 보안 규칙도 "
    "본인 일기 또는 수락된 친구의 공개 일기만 읽을 수 있도록 제한한다. 이는 일반 SNS보다 좁고 안전한 공유 범위를 제공한다.",
)

doc.add_heading("2.7 공유 시트와 PWA", level=2)
add_para(
    doc,
    "일기 상세 화면에는 iOS 스타일 공유 시트가 구현되어 사진 저장, 링크 복사, 메시지 공유와 같은 사용 흐름을 모사한다. "
    "또한 manifest와 service worker가 포함되어 있어 설치형 웹앱(PWA) 경험을 제공할 수 있다.",
)

doc.add_heading("3. 기술 구성", level=1)
add_table(
    doc,
    ["영역", "현재 사용 기술", "역할"],
    [
        ["Frontend", "React 19, Vite, TypeScript, Tailwind CSS, motion/react, lucide-react", "모바일 우선 UI, 탭 네비게이션, 일기 작성·조회 화면 구현"],
        ["Authentication", "Firebase Authentication, Google Provider", "Google 로그인과 사용자 식별"],
        ["Database", "Cloud Firestore", "사용자 프로필, 일기, 친구 요청·관계, 공개 범위 저장"],
        ["Serverless Backend", "Firebase Functions v2, asia-northeast3", "AI API 호출, 입력 검증, 사용량 로그 기록"],
        ["AI Summary", "Gemini 2.5 Flash", "일기 본문을 아이 말투의 짧은 요약문으로 변환"],
        ["AI Image", "Imagen 4.0 Generate", "일기 내용을 크레용 그림일기 이미지로 생성"],
        ["Hosting/PWA", "Firebase Hosting, manifest, service worker", "웹앱 배포, 앱 설치형 경험, 앱 셸 캐싱"],
    ],
    [1750, 3200, 4410],
)

doc.add_heading("3.1 데이터 구조", level=2)
add_table(
    doc,
    ["컬렉션/문서", "주요 필드", "설명"],
    [
        ["users/{email}", "nickname, email, profileColorName, profileColorHex, createdAt, updatedAt", "사용자 프로필과 대표 색상"],
        ["users/{email}/diaries/{date}", "date, title, text, imageUrl, feeling, weather, isPublic, summary, createdAt", "날짜별 그림일기 데이터"],
        ["users/{email}/friends/{friendEmail}", "friendEmail, friendNickname, friendColorHex, status, updatedAt", "친구 요청·수락·삭제 상태"],
        ["aiUsage/{docId}", "ownerEmail, type, createdAt", "AI 요약 및 이미지 생성 사용 로그. 클라이언트 읽기·쓰기는 차단"],
    ],
    [2000, 3900, 3460],
)

doc.add_heading("3.2 AI 처리 흐름", level=2)
add_bullets(
    doc,
    [
        "사용자가 제목, 본문, 날씨, 감정, 공개 여부를 입력한다.",
        "클라이언트가 Firebase callable function으로 일기 본문과 감정 정보를 전달한다.",
        "Gemini 2.5 Flash가 아이 말투의 60자 이내 요약문을 생성한다.",
        "Gemini가 그림 생성을 위한 장면 설명을 짧게 정리한다.",
        "Imagen 4가 크레용 질감의 정사각형 그림일기를 생성한다.",
        "클라이언트는 생성 이미지를 압축해 Firestore 일기 문서에 저장하고 캘린더와 그림책에 표시한다.",
    ],
)

doc.add_heading("4. UI/UX 설계", level=1)
add_para(
    doc,
    "삐뚤의 화면은 스마트폰 세로 화면을 기준으로 설계되어 있다. 앱 전체는 흰색 카드형 모바일 프레임 안에서 동작하며, 하단 탭에는 그림책, 쓰기, 홈, 친구, 설정이 배치된다. "
    "사용자는 복잡한 메뉴를 거치지 않고 일기 쓰기, 월간 기록 보기, 친구 일기 확인으로 바로 이동할 수 있다.",
)
add_para(
    doc,
    "시각 언어는 완성도 높은 일러스트보다 아이가 그린 듯한 서툰 형태와 부드러운 색감을 강조한다. 감정별 색상 카드, 삐뚤한 얼굴 로고, 손그림 느낌의 그림일기 이미지는 "
    "앱 이름 ‘삐뚤’과 일관된 정체성을 만든다.",
)
add_table(
    doc,
    ["화면", "현재 구현 내용"],
    [
        ["스플래시", "삐뚤 로고와 시작 흐름 제공"],
        ["로그인", "Google 계정으로 계속하기"],
        ["프로필 설정", "닉네임 입력, 대표 색상 선택, 커스텀 색상 선택"],
        ["홈/캘린더", "월 이동, 날짜별 일기 카드, 감정 색상 표시"],
        ["쓰기", "제목·본문·날씨·감정·공개 여부 입력, AI 그림 생성"],
        ["일기 상세", "그림, 요약문, 원문, 날씨, 감정, 공개 토글, 공유 시트"],
        ["그림책", "월간 일기를 책처럼 넘겨 보는 화면"],
        ["친구", "친구 목록, 친구 공개 일기, 요청 수락·거절"],
        ["설정", "프로필 수정, 친구 설정, 로그아웃, 계정 삭제"],
    ],
    [2100, 7260],
)

doc.add_heading("5. 개발 현황과 로드맵", level=1)
doc.add_heading("5.1 구현 완료 범위", level=2)
add_bullets(
    doc,
    [
        "React/Vite 기반 모바일 우선 웹앱 화면 구현",
        "Firebase Authentication Google 로그인 구현",
        "Firestore 사용자·일기·친구 데이터 저장 구조 구현",
        "Gemini 요약 및 Imagen 이미지 생성 함수 구현",
        "캘린더, 그림책, 일기 상세, 친구, 설정 화면 구현",
        "친구 요청·수락·거절·삭제와 공개 일기 조회 권한 구현",
        "PWA manifest와 service worker 구성",
    ],
)

doc.add_heading("5.2 향후 개선 과제", level=2)
add_bullets(
    doc,
    [
        "현재 일부 샘플 데이터 문자열의 인코딩 깨짐을 정리하고 실제 한국어 시드 데이터로 교체한다.",
        "Firestore 문서 크기 제한을 고려해 생성 이미지를 Cloud Storage에 저장하고 URL만 Firestore에 보관하는 구조로 확장한다.",
        "공유 시트의 알림 기반 모사 기능을 실제 Web Share API와 다운로드 기능으로 연결한다.",
        "월간 그림책을 PDF 또는 이미지 묶음으로 내보내는 기능을 추가한다.",
        "보호자 모드, 신고·차단, 공개 범위 세분화 등 아동 서비스에 필요한 안전 장치를 강화한다.",
        "Google Calendar 연동은 현재 핵심 기능이 아니므로 선택적 확장 기능으로 재검토한다.",
    ],
)

doc.add_heading("6. 기대 효과", level=1)
doc.add_heading("6.1 일기 작성 부담 완화", level=2)
add_para(
    doc,
    "사용자는 긴 글이나 완성도 높은 그림을 직접 만들 필요 없이, 하루의 일을 짧게 적고 감정과 날씨를 선택하는 것만으로 그림일기를 완성할 수 있다. "
    "AI가 요약과 그림 생성을 담당하기 때문에 기록 진입 장벽이 낮아진다.",
)

doc.add_heading("6.2 회상 경험 강화", level=2)
add_para(
    doc,
    "캘린더와 그림책 화면은 일기를 단순 저장 데이터가 아니라 다시 보고 싶은 시각적 기록으로 만든다. "
    "감정 색상과 날씨 정보가 함께 남아 있어 사용자는 특정 날짜의 분위기와 사건을 빠르게 떠올릴 수 있다.",
)

doc.add_heading("6.3 제한된 관계 안의 공유", level=2)
add_para(
    doc,
    "친구 기능은 공개 일기를 전체 인터넷에 노출하지 않고, 수락된 친구에게만 보여 주는 구조다. "
    "이는 아이 그림일기라는 서비스 성격에 맞게 가볍지만 통제 가능한 공유 경험을 제공한다.",
)

doc.add_heading("6.4 실제 배포 가능한 기술 경험", level=2)
add_para(
    doc,
    "삐뚤은 프론트엔드 화면뿐 아니라 인증, 데이터베이스, 서버리스 AI 호출, 보안 규칙, PWA 구성까지 포함한다. "
    "따라서 단순 프로토타입을 넘어 실제 사용자 흐름을 검증할 수 있는 AI 웹앱 프로젝트로 의미가 있다.",
)

doc.add_heading("7. 결론", level=1)
add_para(
    doc,
    "삐뚤은 초기 제안서의 ‘AI가 일기를 그림으로 바꿔 주는 서비스’라는 핵심 아이디어를 유지하면서, 현재 구현 단계에서는 모바일 우선 그림일기 앱으로 구체화되었다. "
    "사용자는 Google 로그인 후 프로필을 만들고, 하루의 일기를 작성하며, AI가 만든 요약문과 크레용 그림을 캘린더와 그림책으로 회상할 수 있다. "
    "또한 친구 공개 기능을 통해 작은 관계 안에서 서로의 일기를 나눌 수 있다.",
)
add_para(
    doc,
    "따라서 개정된 제안서의 핵심은 Google Calendar나 데스크톱형 복합 기능보다, ‘아이답게 삐뚤한 기록을 쉽고 즐겁게 남기는 AI 그림일기 경험’에 맞춰야 한다. "
    "향후에는 이미지 저장 구조 안정화, 실제 공유 기능, 월간 그림책 내보내기, 보호자 안전 기능을 보강함으로써 서비스 완성도를 높일 수 있다.",
)

footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run("삐뚤 제안서 개정본")
set_run_font(run, size=9, color="555555")

doc.save(OUT)
print(OUT)
