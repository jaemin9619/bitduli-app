from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = r"C:\Users\명재민\Documents\26빅데이터\삐뚤_제안서_기능중심_재작성본.docx"
FONT = "Malgun Gothic"


def set_run(run, size=11, bold=None, color=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for key in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rfonts.set(qn(key), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_para(p, size=11, color=None):
    for run in p.runs:
        set_run(run, size=size, color=color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[idx]))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)


def setup_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    return doc


def para(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    style_para(p)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.194)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)
        style_para(p)


def numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.45)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)
        style_para(p)


def add_table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, header in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = header
        shade(c, "F4F6F9")
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_para(p, size=10, color="0B2545")
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                style_para(p, size=10)
    table_geometry(t, widths)
    doc.add_paragraph()
    return t


doc = setup_doc()

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("삐뚤: AI 그림일기 서비스 제안서")
set_run(r, size=24, bold=True, color="0B2545")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("기능 중심 재작성본")
set_run(r, size=12, color="555555")

doc.add_heading("1. 프로젝트 개요", level=1)
doc.add_heading("1.1 프로젝트명", level=2)
para(doc, "삐뚤")
para(doc, "부제: AI가 도와주는 그림일기 앱")

doc.add_heading("1.2 프로젝트 배경", level=2)
para(
    doc,
    "현대인들은 바쁜 일상 속에서 자신의 하루를 돌아보고 기록할 필요성을 느끼지만, 일기를 꾸준히 작성하는 것은 쉽지 않다. "
    "일기는 감정 정리, 자기 성찰, 일정 회고에 도움을 주는 유용한 기록 수단이지만, 많은 사용자는 글을 길게 작성해야 한다는 부담감, 기록 과정의 단조로움, "
    "지속적인 동기 부족으로 인해 일기 작성을 중단하곤 한다.",
)
para(
    doc,
    "또한 하루를 기록하고 싶어도 무엇을 써야 할지 막막하거나, 작성한 기록이 단순 텍스트로만 남아 다시 확인하고 싶은 흥미를 주지 못하는 경우가 많다. "
    "이에 따라 일기 작성의 진입장벽을 낮추고, 기록 자체를 즐거운 경험으로 전환할 수 있는 새로운 형태의 일기 서비스가 필요하다.",
)

doc.add_heading("1.3 기존 서비스의 한계점", level=2)
para(
    doc,
    "기존 일기 앱은 대부분 텍스트 기록, 사진 첨부, 감정 태그, 캘린더 정리와 같은 기본적인 기록 기능에 초점을 맞추고 있다. "
    "이러한 기능은 사용자의 기록을 저장하고 정리하는 데에는 효과적이지만 다음과 같은 한계가 존재한다.",
)
numbered(
    doc,
    [
        "기록 방식이 단조롭다. 사용자가 직접 텍스트를 입력하고 사진이나 감정 태그를 추가하는 방식에 머물러, 작성 과정에서 새로운 재미나 기대감을 느끼기 어렵다.",
        "사용자 개성을 반영한 시각적 결과물이 부족하다. 기존 서비스는 기록 보관에는 효과적이지만, 사용자의 취향과 감정을 반영한 개인화된 결과물을 제공하는 경우는 제한적이다.",
        "기록과 회상이 분리되어 있다. 날짜별 기록은 저장되지만, 사용자가 자신의 하루를 그림이나 이야기처럼 다시 보고 싶은 경험으로 이어지는 경우가 적다.",
        "지속적인 사용 동기가 부족하다. 일기 작성은 꾸준함이 중요하지만, 작성한 일기가 장기적으로 의미 있는 결과물로 축적되는 보상 경험이 약하다.",
    ],
)

doc.add_heading("1.4 프로젝트 목적", level=2)
para(
    doc,
    "본 프로젝트는 이러한 한계를 해결하기 위해 사용자가 작성한 텍스트 일기를 AI가 분석하고, 하루의 핵심 장면을 그림일기 형태로 변환하는 서비스를 제안한다. "
    "사용자는 일기를 작성하면서 날씨와 감정을 함께 선택하고, AI는 일기 내용에서 사건, 감정, 장소, 행동을 반영한 그림과 짧은 요약문을 생성한다.",
)
para(
    doc,
    "이를 통해 사용자는 단순히 글을 저장하는 것이 아니라, 자신의 하루가 하나의 그림 장면으로 재구성되는 경험을 하게 된다. "
    "삐뚤은 일기 작성의 부담을 줄이고, 기록의 재미와 지속성을 높이며, 개인화된 감성 기록 경험을 제공하는 것을 목표로 한다.",
)

doc.add_heading("1.5 서비스 핵심 콘셉트", level=2)
para(doc, "삐뚤의 핵심 콘셉트는 다음과 같다.")
quote = doc.add_paragraph()
quote.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = quote.add_run("“내 하루가 삐뚤빼뚤한 그림일기가 된다.”")
set_run(r, size=15, bold=True, color="1F4D78")
para(
    doc,
    "삐뚤은 단순한 일기 작성 앱이 아니라, 감정 기록, 날씨 기록, AI 이미지 생성, 캘린더 회상, 월간 그림책, 친구 공유를 결합한 AI 기반 감성 기록 서비스이다. "
    "현재 구현된 기능은 모바일 환경에서 빠르게 일기를 쓰고, AI가 만든 그림과 요약을 통해 다시 보고 싶은 기록으로 축적하는 데 초점을 둔다.",
)

doc.add_heading("2. 주요 기능", level=1)
doc.add_heading("2.1 Google 로그인과 개인 프로필", level=2)
para(
    doc,
    "사용자는 Google 계정으로 로그인한 뒤 닉네임과 대표 색상을 설정한다. 대표 색상은 사용자의 삐뚤 캐릭터와 화면 분위기에 반영되어, 단순 계정 정보가 아니라 개인화된 일기장의 첫 인상을 만든다.",
)

doc.add_heading("2.2 일기 작성 기능", level=2)
para(
    doc,
    "사용자는 날짜를 선택하거나 쓰기 버튼을 눌러 일기를 작성한다. 일기에는 제목, 본문, 날씨, 감정, 공개 여부를 입력할 수 있다. "
    "감정은 행복, 피곤, 평온, 슬픔, 화남 등으로 표현되며, 각 감정은 색상과 표정 요소로 시각화된다.",
)

doc.add_heading("2.3 AI 그림일기 생성", level=2)
para(
    doc,
    "삐뚤의 핵심 기능은 사용자가 작성한 일기를 AI 그림일기로 바꾸는 것이다. AI는 일기 내용을 바탕으로 주요 사건, 분위기, 장소, 행동을 반영한 장면을 생성한다. "
    "현재 서비스의 그림 스타일은 완벽하게 정돈된 일러스트가 아니라, 아이가 크레용으로 그린 듯한 서툰 선과 색 번짐을 강조한다.",
)

doc.add_heading("2.4 AI 요약문 생성", level=2)
para(
    doc,
    "AI는 그림뿐 아니라 일기 내용을 짧은 문장으로 요약한다. 요약문은 보고서식 문장이 아니라 아이가 쓴 그림일기처럼 간단하고 따뜻한 말투를 목표로 한다. "
    "이를 통해 사용자는 긴 본문을 다시 읽지 않아도 그날의 핵심 사건과 감정을 빠르게 떠올릴 수 있다.",
)

doc.add_heading("2.5 캘린더 기반 회상", level=2)
para(
    doc,
    "홈 화면은 월간 캘린더를 중심으로 구성된다. 일기를 작성한 날짜에는 감정 색상과 삐뚤 캐릭터가 표시되어 사용자가 한 달 동안 어떤 감정의 하루를 보냈는지 한눈에 확인할 수 있다. "
    "날짜를 선택하면 해당 일기의 그림, 요약문, 원문, 날씨, 감정, 공개 상태를 확인할 수 있다.",
)

doc.add_heading("2.6 월간 그림책", level=2)
para(
    doc,
    "그림책 화면은 작성한 일기들을 월 단위로 모아 책처럼 넘겨 보는 기능이다. 사용자는 자신의 일기가 단순한 텍스트 목록이 아니라 그림과 짧은 문장으로 구성된 한 권의 그림책처럼 쌓이는 경험을 얻는다. "
    "이는 일기 작성의 지속적인 동기를 만드는 핵심 보상 요소이다.",
)

doc.add_heading("2.7 친구 공유 기능", level=2)
para(
    doc,
    "삐뚤은 개인 기록에 머무르지 않고 친구와 일기를 공유할 수 있는 기능을 제공한다. 사용자는 이메일을 통해 친구를 요청하고, 상대방이 수락하면 서로 친구가 된다. "
    "다만 모든 일기가 자동으로 공개되는 것은 아니며, 사용자가 공개로 설정한 일기만 친구가 볼 수 있다.",
)

doc.add_heading("2.8 설정과 계정 관리", level=2)
para(
    doc,
    "사용자는 설정 화면에서 닉네임과 대표 색상을 수정하고, 친구 요청을 관리하며, 로그아웃 또는 계정 삭제를 진행할 수 있다. "
    "이 기능은 단순한 프로토타입을 넘어 실제 사용자 계정 기반 서비스로 운영될 수 있는 기본 조건을 갖추기 위한 요소이다.",
)

doc.add_heading("3. 사용자 이용 흐름", level=1)
numbered(
    doc,
    [
        "사용자는 앱을 실행하고 Google 계정으로 로그인한다.",
        "처음 사용하는 사용자는 닉네임과 대표 색상을 설정한다.",
        "홈 캘린더에서 오늘 날짜를 선택하거나 쓰기 버튼을 눌러 일기를 작성한다.",
        "제목, 본문, 날씨, 감정, 공개 여부를 입력한다.",
        "AI 그림 생성을 실행하면 그림일기 이미지와 짧은 요약문이 생성된다.",
        "생성 결과를 저장하면 해당 날짜의 일기로 캘린더에 표시된다.",
        "사용자는 캘린더 또는 월간 그림책 화면에서 지난 기록을 다시 본다.",
        "원하는 경우 친구를 추가하고 공개 일기를 서로 확인한다.",
    ],
)

doc.add_heading("4. 기능별 기대 효과", level=1)
add_table(
    doc,
    ["기능", "사용자 경험", "기대 효과"],
    [
        ["AI 그림 생성", "글로 쓴 하루가 그림 장면으로 바뀜", "작성 재미와 재방문 동기 증가"],
        ["AI 요약문", "긴 일기를 짧고 귀여운 문장으로 확인", "회상 부담 감소"],
        ["감정·날씨 기록", "하루의 분위기를 함께 저장", "감정 흐름 파악과 자기 이해 지원"],
        ["캘린더", "날짜별 일기를 한눈에 확인", "기록 누락 확인과 습관 형성"],
        ["월간 그림책", "한 달의 기록을 책처럼 감상", "장기적인 축적 보상 제공"],
        ["친구 공유", "선택한 일기만 친구에게 공개", "부담 없는 소셜 기록 경험"],
        ["프로필 색상", "내 일기장의 시각적 정체성 형성", "개인화 경험 강화"],
    ],
    [2100, 3650, 3610],
)

doc.add_heading("5. 서비스 차별점", level=1)
doc.add_heading("5.1 기록을 결과물로 바꾸는 경험", level=2)
para(
    doc,
    "삐뚤은 사용자가 입력한 글을 단순히 저장하는 데 그치지 않고, AI가 그림과 요약문이라는 결과물로 바꾸어 준다. "
    "사용자는 일기를 작성할 때마다 어떤 그림이 나올지 기대하게 되며, 이는 기존 텍스트 중심 일기 앱과 구별되는 핵심 차별점이다.",
)

doc.add_heading("5.2 아이답고 서툰 시각 스타일", level=2)
para(
    doc,
    "일반적인 AI 이미지는 지나치게 정교하거나 완성도가 높아 일기라는 개인 기록의 감성과 어울리지 않을 수 있다. "
    "삐뚤은 삐뚤한 선, 크레용 질감, 선 밖으로 번지는 색감처럼 아이 그림일기의 특징을 강조하여 더 따뜻하고 친근한 기록 경험을 제공한다.",
)

doc.add_heading("5.3 캘린더와 그림책의 결합", level=2)
para(
    doc,
    "캘린더는 기록의 위치를 알려 주고, 그림책은 기록을 감상하게 만든다. 삐뚤은 이 두 화면을 함께 제공하여 사용자가 일기를 관리하는 동시에 자신의 시간을 이야기처럼 되돌아볼 수 있게 한다.",
)

doc.add_heading("5.4 제한된 공유 구조", level=2)
para(
    doc,
    "삐뚤의 친구 기능은 일반 SNS처럼 무작위로 공개되는 구조가 아니라, 친구 요청과 수락을 기반으로 한다. "
    "또한 사용자가 공개로 설정한 일기만 친구에게 보여 주기 때문에 개인 기록의 성격을 유지하면서도 가벼운 소통을 가능하게 한다.",
)

doc.add_heading("6. 현재 구현된 핵심 기능과 향후 발전점", level=1)
doc.add_heading("6.1 현재 구현된 핵심 기능", level=2)
bullets(
    doc,
    [
        "Google 계정 로그인",
        "닉네임과 대표 색상 기반 프로필 설정",
        "날짜별 일기 작성",
        "날씨와 감정 선택",
        "AI 그림일기 이미지 생성",
        "AI 요약문 생성",
        "월간 캘린더 기반 기록 확인",
        "월간 그림책 보기",
        "친구 요청, 승인, 거절, 삭제",
        "공개 일기만 친구에게 공유",
        "프로필 수정, 로그아웃, 계정 삭제",
        "모바일 우선 PWA 형태의 앱 경험",
    ],
)

doc.add_heading("6.2 향후 발전점", level=2)
para(
    doc,
    "다음 기능들은 서비스의 확장 가능성을 보여 주는 발전 방향이다. 현재 제안서의 본문에서는 구현된 핵심 기능을 중심으로 설명하고, 아래 항목은 향후 고도화 과제로 분리한다.",
)
bullets(
    doc,
    [
        "캐릭터 커스터마이징 고도화: 현재는 닉네임과 대표 색상 중심이지만, 향후 머리 모양, 의상, 표정, 액세서리 등으로 확장할 수 있다.",
        "라인 아트 스타일 선택: 현재는 크레용 그림일기 스타일을 중심으로 하며, 향후 라인 아트, 색연필, 수채화 등 스타일 선택 기능을 제공할 수 있다.",
        "외부 캘린더 연동: 현재는 앱 내부 캘린더를 사용하며, 향후 Google Calendar 등 외부 일정과 연동해 일정 회고 기능을 강화할 수 있다.",
        "월간 그림책 발행: 현재는 앱 안에서 월간 그림책을 볼 수 있으며, 향후 PDF, 이미지 묶음, 인쇄용 책자 형태로 내보내는 기능을 추가할 수 있다.",
        "공유 기능 고도화: 현재는 친구 공개 중심이며, 향후 이미지 다운로드, Web Share API, 가족 앨범 공유 등으로 확장할 수 있다.",
        "보호자 및 안전 기능: 아동·가족 사용자를 고려해 보호자 확인, 신고, 차단, 공개 범위 세분화 기능을 추가할 수 있다.",
        "기록 분석 리포트: 감정, 날씨, 키워드 흐름을 월간 리포트로 제공해 자기 이해와 회고 경험을 강화할 수 있다.",
    ],
)

doc.add_heading("7. 기대 효과", level=1)
doc.add_heading("7.1 일기 작성 부담 완화", level=2)
para(
    doc,
    "사용자는 긴 글을 완성해야 한다는 부담 없이 하루의 일을 간단히 적고 감정과 날씨를 선택하는 것만으로 그림일기를 만들 수 있다. "
    "AI가 그림과 요약을 보조하기 때문에 일기 작성의 진입장벽이 낮아진다.",
)

doc.add_heading("7.2 기록의 재미와 지속성 강화", level=2)
para(
    doc,
    "일기를 작성할 때마다 새로운 그림이 생성되고, 그 결과가 캘린더와 그림책에 축적된다. 사용자는 기록이 쌓이는 과정을 눈으로 확인하며 지속적으로 앱을 사용할 동기를 얻는다.",
)

doc.add_heading("7.3 개인화된 감성 기록 제공", level=2)
para(
    doc,
    "대표 색상, 감정, 날씨, AI 그림이 결합되어 사용자의 하루가 개인화된 시각적 기록으로 남는다. 이는 단순 텍스트 저장보다 더 강한 정서적 몰입감을 제공한다.",
)

doc.add_heading("7.4 관계 기반 회상 경험 제공", level=2)
para(
    doc,
    "친구 공개 기능을 통해 사용자는 자신의 일부 기록을 가까운 사람과 나눌 수 있다. 단, 공개 여부를 사용자가 직접 선택하기 때문에 개인 기록의 안전성과 공유의 즐거움을 함께 확보할 수 있다.",
)

doc.add_heading("8. 결론", level=1)
para(
    doc,
    "삐뚤은 일기 작성의 부담을 낮추고, 사용자의 하루를 AI 그림일기로 재구성하는 감성 기록 서비스이다. "
    "현재 구현된 핵심 기능은 Google 로그인, 프로필 설정, 일기 작성, 감정·날씨 기록, AI 그림 생성, AI 요약, 캘린더 회상, 월간 그림책, 친구 공유로 구성된다.",
)
para(
    doc,
    "본 서비스는 기존 일기 앱의 단조로운 기록 방식을 보완하고, 사용자가 자신의 하루를 다시 보고 싶은 그림 장면으로 남기게 한다. "
    "향후 캐릭터 커스터마이징, 외부 캘린더 연동, 그림책 발행, 보호자 안전 기능 등을 확장한다면 AI 기반 감성 기록 플랫폼으로 발전할 수 있다.",
)

footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run("삐뚤 제안서 기능 중심 재작성본")
set_run(r, size=9, color="555555")

doc.save(OUT)
print(OUT)
